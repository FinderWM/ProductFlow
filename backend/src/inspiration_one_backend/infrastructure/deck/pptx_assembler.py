from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.util import Emu

# 标准 16:9 宽屏画布（EMU）：13.333in × 7.5in。
_SLIDE_WIDTH_EMU = 12192000
_SLIDE_HEIGHT_EMU = 6858000
_BLANK_LAYOUT_INDEX = 6


def build_deck_pptx(slides: list[tuple[bytes, str | None]]) -> bytes:
    """把每页整页图满铺为一张 16:9 幻灯片，演讲备注写入备注栏，返回 pptx 字节。"""
    presentation = Presentation()
    presentation.slide_width = Emu(_SLIDE_WIDTH_EMU)
    presentation.slide_height = Emu(_SLIDE_HEIGHT_EMU)
    blank_layout = presentation.slide_layouts[_BLANK_LAYOUT_INDEX]
    for image_bytes, notes in slides:
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            BytesIO(image_bytes),
            Emu(0),
            Emu(0),
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
